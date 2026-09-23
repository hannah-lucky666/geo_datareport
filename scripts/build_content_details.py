#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
从「投放文章统计」Excel 生成 src/data/content_details.json，供投放明细页取数。

用法:
  python scripts/build_content_details.py
  python scripts/build_content_details.py --toc <ToC.xlsx> --tob <ToB.xlsx>
"""

import argparse
import io
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_DIR = r"f:\9.16\02GEO\3.GEO数据监测系统\9.21投放文章汇总\投放分析 0921\投放分析 0921"

PLATFORM_COLS = [
    ("deepseek", "DeepSeek 引用次数"),
    ("doubao", "豆包 引用次数"),
    ("qwen", "通义千问 引用次数"),
]

SUMMARY_MARKERS = ("【汇总】", "汇总", "总计", "合计")


def to_int(v):
    if v is None or v == "":
        return 0
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def read_overview(ws):
    wanted = ("总会话数", "引用记录总数", "投放文章引用次数", "投放文章总数", "被引用文章数")
    out = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() in wanted:
                for dr in (1, 2):
                    cand = ws.cell(r + dr, c).value
                    if cand not in (None, ""):
                        out[v.strip()] = to_int(cand)
                        break
    return out


def normalize_channel(name):
    return re.sub(r"[（(].*", "", name).strip()


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    overview = read_overview(wb["数据概览"])

    ws = wb["文章引用明细"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {name: header.index(name) for name in header if name}
    platforms = [(key, label) for key, label in PLATFORM_COLS if label in col]

    articles = []
    summary_row = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = row[col["文章标题"]]
        if title in (None, ""):
            continue
        title = str(title).strip()
        if any(m in title for m in SUMMARY_MARKERS):
            summary_row = row
            continue
        articles.append({
            "rank": to_int(row[col["排名"]]),
            "title": title,
            "platform": str(row[col["发布平台"]] or "").strip(),
            "domain": str(row[col["链接域名"]] or "").strip(),
            "date": str(row[col["发布时间"]] or "").strip()[:10],
            "total": to_int(row[col["总引用次数"]]),
            "isCited": str(row[col["是否被引用"]] or "").strip(),
            "link": str(row[col["文章链接"]] or "").strip(),
            **{key: to_int(row[col[label]]) for key, label in platforms},
        })

    by_platform = {key: sum(a[key] for a in articles) for key, _ in platforms}
    total_citations = sum(a["total"] for a in articles)
    cited_articles = sum(1 for a in articles if a["isCited"] == "是")

    checks = []
    if summary_row is not None:
        exp_total = to_int(summary_row[col["总引用次数"]])
        if exp_total != total_citations:
            checks.append(f"总引用次数 明细求和={total_citations} 与汇总行={exp_total} 不一致")
        for key, label in platforms:
            exp = to_int(summary_row[col[label]])
            if exp != by_platform[key]:
                checks.append(f"{label} 明细求和={by_platform[key]} 与汇总行={exp} 不一致")
    if overview.get("投放文章引用次数") not in (None, total_citations):
        checks.append(f"总引用次数 明细求和={total_citations} 与概览页={overview.get('投放文章引用次数')} 不一致")
    if overview.get("被引用文章数") not in (None, cited_articles):
        checks.append(f"被引用文章数 明细={cited_articles} 与概览页={overview.get('被引用文章数')} 不一致")

    ch = wb["平台与渠道分析"]
    raw_channels = []
    for row in ch.iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            continue
        raw_channels.append({"name": str(row[0]).strip(), "citations": to_int(row[4])})
    raw_channels.sort(key=lambda x: -x["citations"])
    channels, seen = [], set()
    for c in raw_channels:
        if c["citations"] <= 0:
            continue
        n = normalize_channel(c["name"])
        if n and n not in seen and not re.fullmatch(r"[a-z0-9.\-]+", n):
            seen.add(n)
            channels.append(n)

    total_articles = overview.get("投放文章总数") or len(articles)
    return {
        "total_articles": total_articles,
        "cited_articles": cited_articles,
        "citation_rate": round(cited_articles / total_articles * 100, 1) if total_articles else 0,
        "total_citations": total_citations,
        "by_platform": by_platform,
        "channels": channels[:6],
        "top10": articles[:10],
        "article_rows": len(articles),
        "checks": checks,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--toc", default=os.path.join(XLSX_DIR, "桃李面包-优化词-ToC_投放文章统计_0921.xlsx"))
    ap.add_argument("--tob", default=os.path.join(XLSX_DIR, "桃李面包-优化词-ToB_投放文章统计_0921.xlsx"))
    args = ap.parse_args()

    targets = [
        ("toc", "桃李面包 ToC", args.toc),
        ("tob", "桃李面包 ToB", args.tob),
    ]
    out = {"meta": {"generated_at": datetime.now().isoformat(timespec="seconds"), "sources": {}}}

    ok = True
    for key, name, path in targets:
        if not os.path.exists(path):
            print(f"找不到文件: {path}")
            ok = False
            continue
        data = parse_workbook(path)
        data["product"] = name
        out["meta"]["sources"][key] = os.path.basename(path)
        checks = data.pop("checks")
        out[key] = data

        print(f"=== {name} ===")
        print(
            f'  投放文章 {data["total_articles"]} 篇（明细 {data["article_rows"]} 行）'
            f' / 被引 {data["cited_articles"]} 篇 = {data["citation_rate"]}%'
        )
        print(
            f'  引用总次数 {data["total_citations"]}  分平台 {data["by_platform"]}'
            f'  校验和 {sum(data["by_platform"].values())}'
        )
        print(f'  覆盖渠道 {"、".join(data["channels"])}')
        if data["top10"]:
            print(f'  Top10 引用区间 {data["top10"][0]["total"]} ~ {data["top10"][-1]["total"]}')
        if checks:
            ok = False
            for c in checks:
                print(f"  !! 校验失败: {c}")
        else:
            print("  校验通过：明细求和 = 汇总行 = 概览页")

    out_path = os.path.join(ROOT, "src", "data", "content_details.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"\n已写入 {out_path}")
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
